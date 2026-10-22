"""Word 公式抽取（pandoc OMML→LaTeX）测试。"""
import subprocess
from pathlib import Path

from docx import Document

import app.core.ingest.docx as docx_mod


def _save_minimal_docx(path):
    d = Document()
    d.add_paragraph("x")
    d.save(str(path))


def test_omml_docx_uses_pandoc(tmp_path, monkeypatch):
    path = tmp_path / "f.docx"
    _save_minimal_docx(path)
    monkeypatch.setattr(docx_mod, "_has_omml", lambda doc: True)
    monkeypatch.setattr(
        docx_mod, "_pandoc_to_markdown", lambda p: "# 标题\n\n公式 $E=mc^2$\n\n正文"
    )
    parsed = docx_mod.parse_docx(path)
    assert "$E=mc^2$" in parsed.full_text
    assert parsed.engine_used == "pandoc"
    assert parsed.title == "标题"
    assert any(b.type == "heading" for b in parsed.blocks)
    assert any(b.type == "paragraph" and "正文" in b.text for b in parsed.blocks)


def test_no_omml_uses_python_docx(tmp_path, monkeypatch):
    path = tmp_path / "f.docx"
    _save_minimal_docx(path)
    monkeypatch.setattr(docx_mod, "_has_omml", lambda doc: False)
    called = []
    monkeypatch.setattr(docx_mod, "_pandoc_to_markdown", lambda p: called.append(True))
    parsed = docx_mod.parse_docx(path)
    assert not called  # 未走 pandoc
    assert parsed.engine_used == "python-docx"


def test_pandoc_missing_falls_back_to_python_docx(tmp_path, monkeypatch):
    """pandoc 缺失时回退 python-docx 结构化解析，非公式正文不丢失。"""
    path = tmp_path / "f.docx"
    d = Document()
    d.add_paragraph("保留这段正文")
    d.save(str(path))

    monkeypatch.setattr(docx_mod, "_has_omml", lambda doc: True)
    monkeypatch.setattr(
        docx_mod, "_pandoc_to_markdown", lambda p: (_ for _ in ()).throw(FileNotFoundError())
    )
    parsed = docx_mod.parse_docx(path)
    assert parsed.engine_used == "python-docx"
    assert "保留这段正文" in parsed.full_text


def test_pandoc_to_markdown_decodes_utf8(monkeypatch):
    """pandoc 输出为 UTF-8，Windows 下 text=True 需显式指定 encoding 避免 gbk 解码崩溃。"""
    captured = {}

    class FakeResult:
        returncode = 0
        stdout = "# 标题\n\n$E=mc^2$\n"
        stderr = ""

    monkeypatch.setattr(
        subprocess, "run", lambda cmd, **kwargs: captured.update(kwargs) or FakeResult()
    )
    out = docx_mod._pandoc_to_markdown(Path("f.docx"))
    assert captured.get("encoding") == "utf-8"
    assert "$E=mc^2$" in out
