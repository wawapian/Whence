import time
from pathlib import Path

import docx as _docx
import pymupdf as fitz
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app import deps
from app.api import files as files_mod
from app.config import Settings
from app.core import llm
from app.core.ingest import ocr, pdf
from app.core.ingest.models import ImageRef, ParsedDocument
from app.main import create_app
from app.models.db import Image


def test_parsed_document_images_defaults_empty():
    doc = ParsedDocument(
        kind="word",
        title=None,
        full_text="",
        headings=[],
        blocks=[],
        engine_used="x",
    )
    assert doc.images == []


def test_imageref_defaults():
    ref = ImageRef(blob=b"\x00")
    assert ref.page is None
    assert ref.position is None
    assert ref.ocr_text == ""


class _FakeOcr:
    def recognize(self, image: bytes) -> str:
        return "查重报告 相似度15%"


def _make_docx_with_image(tmp_path) -> Path:
    png = tmp_path / "r.png"
    PILImage.new("RGB", (1, 1), "white").save(png)
    d = _docx.Document()
    d.add_paragraph("这是正文")
    d.add_picture(str(png))
    p = tmp_path / "f.docx"
    d.save(str(p))
    return p


def _client(tmp_path, monkeypatch):
    app = create_app()
    app.dependency_overrides[deps.get_settings] = lambda: Settings(
        data_dir=tmp_path, embedding_model="BAAI/bge-small-zh-v1.5")
    monkeypatch.setattr(llm, "embed", lambda text, model=None: [0.1] * 512)
    monkeypatch.setattr(ocr, "get_ocr_engine", lambda name: _FakeOcr())
    return TestClient(app)


def _upload_docx(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    path = _make_docx_with_image(tmp_path)
    with open(path, "rb") as f:
        files = {"file": ("f.docx", f,
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        resp = client.post("/api/files", files=files,
                           data={"allow_cloud": "0", "multimodal": "0", "ocr": "auto"})
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]
    for _ in range(50):
        time.sleep(0.1)
        d = client.get(f"/api/files/{doc_id}").json()
        if d["status"] == "done":
            break
    assert d["status"] == "done"
    return client, doc_id


def test_upload_docx_persists_images(tmp_path, monkeypatch):
    client, doc_id = _upload_docx(tmp_path, monkeypatch)

    session = Session(create_engine(f"sqlite:///{tmp_path}/mem.db"))
    try:
        rows = session.query(Image).filter(Image.document_id == doc_id).all()
        assert len(rows) == 1
        assert rows[0].ocr_text == "查重报告 相似度15%"
        assert rows[0].page is None
        assert rows[0].file_path is not None
        disk_path = tmp_path / rows[0].file_path
        assert disk_path.exists()
        assert disk_path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
    finally:
        session.close()


def test_image_disk_save_failure_does_not_fail_document(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)

    def boom(data_dir, blob):
        raise OSError("disk full")

    monkeypatch.setattr(files_mod, "_save_image_blob", boom)

    path = _make_docx_with_image(tmp_path)
    with open(path, "rb") as f:
        files = {"file": ("f.docx", f,
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        resp = client.post("/api/files", files=files,
                           data={"allow_cloud": "0", "multimodal": "0", "ocr": "auto"})
    doc_id = resp.json()["document_id"]
    for _ in range(50):
        time.sleep(0.1)
        d = client.get(f"/api/files/{doc_id}").json()
        if d["status"] == "done":
            break
    assert d["status"] == "done"

    session = Session(create_engine(f"sqlite:///{tmp_path}/mem.db"))
    try:
        rows = session.query(Image).filter(Image.document_id == doc_id).all()
        assert len(rows) == 1
        assert rows[0].file_path is None
        assert rows[0].ocr_text == "查重报告 相似度15%"
    finally:
        session.close()


def test_pdf_parse_collects_images(tmp_path, monkeypatch):
    png = tmp_path / "r.png"
    PILImage.new("RGB", (1, 1), "white").save(png)
    p = tmp_path / "f.pdf"
    d = fitz.open()
    page = d.new_page()
    page.insert_text((72, 72), "正文", fontname="china-s")
    page.insert_image(fitz.Rect(72, 100, 200, 200), filename=str(png))
    d.save(str(p))
    d.close()

    monkeypatch.setattr(ocr, "get_ocr_engine", lambda name: _FakeOcr())
    parsed = pdf.parse_pdf_with_images(p, ocr_name="rapidocr")
    assert len(parsed.images) == 1
    assert parsed.images[0].page == 1
    assert parsed.images[0].ocr_text == "查重报告 相似度15%"
