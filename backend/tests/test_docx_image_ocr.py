from pathlib import Path

import docx as _docx
from PIL import Image

from app.core.ingest import docx as docx_mod, ocr


class _FakeOcr:
    def recognize(self, image: bytes) -> str:
        return "查重报告 相似度15%"


def _make_docx_with_image(tmp_path) -> Path:
    png = tmp_path / "r.png"
    Image.new("RGB", (1, 1), "white").save(png)
    d = _docx.Document()
    d.add_picture(str(png))
    p = tmp_path / "f.docx"
    d.save(str(p))
    return p


def test_image_ocr_text_in_block(tmp_path, monkeypatch):
    monkeypatch.setattr(ocr, "get_ocr_engine", lambda name: _FakeOcr())
    parsed = docx_mod.parse_docx(_make_docx_with_image(tmp_path), ocr_name="rapidocr")
    assert "查重报告" in parsed.full_text
    assert any(b.type == "image" and "查重报告" in b.text for b in parsed.blocks)
    assert parsed.ocr_used is True


def test_ocr_off_leaves_image_empty(tmp_path):
    parsed = docx_mod.parse_docx(_make_docx_with_image(tmp_path), ocr_name="off")
    assert any(b.type == "image" and b.text == "" for b in parsed.blocks)
    assert "查重报告" not in parsed.full_text
    assert parsed.ocr_used is False


def test_ocr_failure_skips_image(tmp_path, monkeypatch):
    def boom(image):
        raise RuntimeError("ocr fail")

    monkeypatch.setattr(
        ocr, "get_ocr_engine", lambda name: type("X", (), {"recognize": boom})()
    )
    parsed = docx_mod.parse_docx(_make_docx_with_image(tmp_path), ocr_name="rapidocr")
    assert any(b.type == "image" and b.text == "" for b in parsed.blocks)
    assert "查重报告" not in parsed.full_text


def test_multi_image_one_fails_keeps_other(tmp_path, monkeypatch):
    # 同一段落内两张图：第一张 OCR 成功、第二张抛异常，应保留第一张文字。
    png1 = tmp_path / "a.png"
    png2 = tmp_path / "b.png"
    Image.new("RGB", (1, 1), "white").save(png1)
    Image.new("RGB", (1, 1), "black").save(png2)
    d = _docx.Document()
    p = d.add_paragraph()
    p.add_run().add_picture(str(png1))
    p.add_run().add_picture(str(png2))
    f = tmp_path / "two.docx"
    d.save(str(f))

    class _FlakyOcr:
        def __init__(self):
            self.calls = 0

        def recognize(self, image: bytes) -> str:
            self.calls += 1
            if self.calls == 1:
                return "图一文字"
            raise RuntimeError("fail second")

    monkeypatch.setattr(ocr, "get_ocr_engine", lambda name: _FlakyOcr())
    parsed = docx_mod.parse_docx(f, ocr_name="rapidocr")
    assert "图一文字" in parsed.full_text
    assert parsed.ocr_used is True


def test_blob_extraction_failure_skips_image(tmp_path, monkeypatch):
    # _image_blobs 抛异常（关系损坏）应跳过该图，不使整篇文档失败，其余正文仍解析。
    png = tmp_path / "r.png"
    Image.new("RGB", (1, 1), "white").save(png)
    d = _docx.Document()
    d.add_paragraph("正文段落")
    d.add_picture(str(png))
    f = tmp_path / "broken_blob.docx"
    d.save(str(f))

    def boom(paragraph):
        raise KeyError("broken relationship")

    monkeypatch.setattr(docx_mod, "_image_blobs", boom)
    monkeypatch.setattr(ocr, "get_ocr_engine", lambda name: _FakeOcr())
    parsed = docx_mod.parse_docx(f, ocr_name="rapidocr")
    assert "正文段落" in parsed.full_text
    assert any(b.type == "image" and b.text == "" for b in parsed.blocks)
    assert parsed.ocr_used is False
