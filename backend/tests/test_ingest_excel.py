"""Excel（.xlsx/.xlsm）支持：Docling 后端解析 + 结构化表 + 全链路 ingest。

试金石后的第一个新格式（只加类 + 注册，未动 detect/report/pipeline）。
"""
from pathlib import Path

from openpyxl import Workbook

from app.core.ingest import pipeline, processors
from app.core.ingest.excel_docling import parse_excel_docling
from app.core.ingest.report import DataReport


def _make_xlsx(tmp_path) -> Path:
    """2 个 sheet：合并标题 + 课程表 + 单格说明（单格应转文本而非 1×1 表）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩总表"
    ws.merge_cells("A1:C1")
    ws["A1"] = "2026 学年成绩"
    ws.append(["课程", "学分", "成绩"])
    ws.append(["体育", 0.5, 88])
    ws.append(["高数", 3.0, 75])
    ws2 = wb.create_sheet("说明")
    ws2["A1"] = "本表由教务处提供"
    path = tmp_path / "成绩.xlsx"
    wb.save(path)
    return path


def _report(suffix: str) -> DataReport:
    return DataReport(kind="other", suffix=suffix, has_text_layer=None)


def test_parse_extracts_text_and_tables(tmp_path):
    parsed = parse_excel_docling(_make_xlsx(tmp_path))
    assert parsed.engine_used == "docling-xlsx"
    assert "2026 学年成绩" in parsed.full_text          # 合并标题成文本
    assert "本表由教务处提供" in parsed.full_text        # 单格默认转文本（非 1×1 假表）
    assert len(parsed.tables) == 1
    table = parsed.tables[0]
    assert table.headers == ["课程", "学分", "成绩"]
    assert [r[0] for r in table.rows] == ["体育", "高数"]
    assert [r[1] for r in table.rows] == ["0.5", "3"]   # 数值原样（归一化不臆改）


def _make_titled_xlsx(tmp_path) -> Path:
    """标题格未合并（A1 打字、B1 起为空）+ 第二行真表头 —— 常见真实形态。"""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "三年级四班（东校区）期末成绩表"
    ws.append(["姓名", "语文", "数学", "英语"])
    ws.append(["甲", 90, 80, 70])
    ws.append(["乙", 60, 95, 88])
    path = tmp_path / "grades.xlsx"
    wb.save(path)
    return path


def test_unmerged_title_row_is_not_header(tmp_path):
    """未合并的标题行不得吞掉真表头：表头=第二行、标题进 heading、表不再 suspect。"""
    parsed = parse_excel_docling(_make_titled_xlsx(tmp_path))
    assert len(parsed.tables) == 1
    table = parsed.tables[0]
    assert table.headers == ["姓名", "语文", "数学", "英语"]
    assert [r[0] for r in table.rows] == ["甲", "乙"]
    assert table.suspect is False
    assert table.heading == "三年级四班（东校区）期末成绩表"


def test_processor_claims_xlsx_only(tmp_path):
    excel = processors.ExcelProcessor()
    assert excel.applies(_report(".xlsx"))
    assert excel.applies(_report(".xlsm"))
    assert not excel.applies(_report(".xls"))           # 老二进制格式不收，如实标注
    assert not excel.applies(_report(".docx"))
    chain = processors.build_chain(_report(".xlsx"))
    assert [type(p).__name__ for p in chain] == ["ExcelProcessor"]
    fallback = processors.build_chain(_report(".xls"))
    assert isinstance(fallback[0], processors.UnknownProcessor)


def test_pipeline_ingest_xlsx_full_chain(tmp_path):
    result = pipeline.ingest(_make_xlsx(tmp_path), ocr_name="off")
    assert result.parsed.engine_used == "docling-xlsx"
    assert result.tables and result.tables[0].headers == ["课程", "学分", "成绩"]
    # 表格行组块进了分块列表（后续走既有 FTS/向量索引 → 普通问答与 table_stats 都能用）
    assert any("课程" in c.text for c in result.chunks)
    # 配对成立：full_text 里的 markdown 管道表原文被替换为表格摘要（表格原文不进正文索引）
    assert "表格（" in result.parsed.full_text


def test_merged_company_cell_fills_down(tmp_path):
    """纵向合并格（企业名跨多行）向下填充到覆盖行——否则续行的岗位失去归属。

    真实案例：岗位信息表一家公司 N 个岗位共用一个合并的公司名格，
    不填充则「这家公司有哪些岗位」只能查到第一条。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["企业名称", "招聘岗位", "人数"])
    ws.merge_cells("A2:A4")
    ws["A2"] = "示例市样例电子有限公司"
    ws["B2"], ws["C2"] = "技术工程师", 6
    ws["B3"], ws["C3"] = "软件工程师", 2
    ws["B4"], ws["C4"] = "PQC工程师", 6
    path = tmp_path / "岗位.xlsx"
    wb.save(path)

    parsed = parse_excel_docling(path)
    table = parsed.tables[0]
    assert [row[0] for row in table.rows] == ["示例市样例电子有限公司"] * 3  # 合并格向下填充
    assert [row[1] for row in table.rows] == ["技术工程师", "软件工程师", "PQC工程师"]


def test_unsupported_xls_stays_honest(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"fake-xls-data")
    result = pipeline.ingest(path, ocr_name="off")
    assert result.parsed.engine_used == "unsupported"    # 不崩、不假装读到了
    assert result.chunks == []
