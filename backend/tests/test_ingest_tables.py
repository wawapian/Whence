from app.core.ingest import pipeline as pl
from app.core.ingest import processors as proc
from app.core.ingest.models import Block, ParsedDocument
from app.core.tables import ParsedTable


def _docling_fake():
    md_table = "|A|B|\n|---|---|\n|1|2|"
    return ParsedDocument(
        kind="pdf", title=None, full_text="# 表\n\n" + md_table, headings=[],
        blocks=[Block(type="table", text=md_table, heading_path="表")],
        engine_used="docling",
    )


def test_pdf_text_layer_uses_coordinate_tables(tmp_path, monkeypatch):
    from helpers_pdf import make_table_pdf
    p = tmp_path / "t.pdf"; make_table_pdf(p)
    monkeypatch.setattr(proc, "parse_pdf_docling", lambda _p: _docling_fake())
    parsed = proc._parse_pdf(p, "off", "docling")
    assert parsed.tables and parsed.tables[0].headers == ["course", "credit"]   # 坐标法产出
    block = [b for b in parsed.blocks if b.type == "table"][0]
    assert block.text.startswith("表格（")                                      # docling 表块换摘要


def test_pdf_coordinate_zero_output_keeps_docling_tables(tmp_path, monkeypatch):
    from helpers_pdf import make_table_pdf
    p = tmp_path / "t.pdf"; make_table_pdf(p)
    monkeypatch.setattr(proc, "parse_pdf_docling", lambda _p: _docling_fake())
    monkeypatch.setattr(proc.pdf_tables, "extract_pdf_tables", lambda _p: [])
    parsed = proc._parse_pdf(p, "off", "docling")
    assert parsed.tables == []
    assert [b for b in parsed.blocks if b.type == "table"][0].text.startswith("|A|B|")  # 保留


def test_scanned_pdf_skips_coordinate_path(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(proc.pdf, "parse_pdf_pages", lambda _p: [""])          # 无文字层
    monkeypatch.setattr(proc.pdf_tables, "extract_pdf_tables",
                        lambda _p: calls.append(1) or [])
    monkeypatch.setattr(proc, "parse_pdf_docling", lambda _p: _docling_fake())
    parsed = proc._parse_pdf(tmp_path / "s.pdf", "off", "docling")
    assert calls == [] and parsed.tables == []


def test_lightweight_pdf_excludes_table_text_from_full_text(tmp_path):
    from helpers_pdf import make_table_pdf
    p = tmp_path / "t.pdf"; make_table_pdf(p)
    parsed = proc._parse_pdf(p, "off", "off")            # 轻量路径（docling 关）
    assert parsed.tables and parsed.tables[0].rows == [["PE", "1.0"]]
    assert "PE" not in parsed.full_text and "course" not in parsed.full_text   # 表格区域文字不进正文
    assert parsed.needs_ocr is False and parsed.ocr_used is False               # 判定仍按未过滤文本


def test_table_page_not_mistaken_for_scanned_page(tmp_path, monkeypatch):
    # 混合文档（表格页 + 空页）开了 OCR：表格区域过滤后正文为空，但该页有文字层，不算扫描页
    import pymupdf
    from helpers_pdf import make_table_pdf
    p1 = tmp_path / "t.pdf"; make_table_pdf(p1)
    doc = pymupdf.open(str(p1)); doc.new_page(); p = tmp_path / "mixed.pdf"
    doc.save(str(p)); doc.close()

    calls = []

    class FakeEngine:
        def recognize(self, image: bytes) -> str:
            calls.append(1)
            return "扫描页文字"

    monkeypatch.setattr(proc.ocr, "get_ocr_engine", lambda name: FakeEngine())
    parsed = proc._parse_pdf(p, "rapidocr", "off")
    assert len(calls) == 1                      # 只 OCR 真正无文字层的页
    assert "扫描页文字" in parsed.full_text
    assert "course" not in parsed.full_text and "PE" not in parsed.full_text
    assert parsed.ocr_used is True and parsed.tables


def test_structify_prefers_structured_and_summarizes_blocks():
    parsed = ParsedDocument(
        kind="pdf", title=None, full_text="# 表\n\n|A|B|\n|---|---|\n|1|2|", headings=[],
        blocks=[Block(type="table", text="|A|B|\n|---|---|\n|1|2|", heading_path="表")],
        engine_used="t",
        tables=[ParsedTable(headers=["课程"], rows=[["体育"]])],
    )
    ts = pl.structify_tables(parsed)
    assert ts == parsed.tables
    assert parsed.blocks[0].text.startswith("表格（1行×1列")      # 块文本替换为摘要
    assert "|A|B|" not in parsed.full_text                       # full_text 同步替换


def test_structify_syncs_full_text_when_block_already_summarized():
    # pandoc 路径：来源已把块文本换成摘要、full_text 仍含原表 → full_text 也要替换（Ruling ⚠️②）
    raw = "|A|B|\n|---|---|\n|1|2|"
    parsed = ParsedDocument(
        kind="word", title=None, full_text="# 标题\n\n" + raw, headings=[],
        blocks=[Block(type="table", text="表格（1行×2列：A、B）", heading_path="标题")],
        engine_used="pandoc",
        tables=[ParsedTable(headers=["A", "B"], rows=[["1", "2"]])],
    )
    pl.structify_tables(parsed)
    assert "|A|B|" not in parsed.full_text
    assert "表格（1行×2列" in parsed.full_text


def test_structify_falls_back_to_markdown_block():
    parsed = ParsedDocument(
        kind="text", title=None, full_text="|a|b|\n|---|---|\n|1|2|", headings=[],
        blocks=[Block(type="table", text="|a|b|\n|---|---|\n|1|2|")], engine_used="t",
    )
    ts = pl.structify_tables(parsed)
    assert len(ts) == 1 and ts[0].headers == ["a", "b"]


def test_structify_no_mislabeling_when_pandoc_subtables_split(tmp_path, monkeypatch):
    # pandoc 路径 + 并排子表：一张 docx 表被切成 2 张结构化表，full_text 管道表数(2) ≠ 结构化表数(3)
    # → 不得按序错配（否则第二张表的原文会被盖上「子表 2」的摘要）
    import docx as D
    from app.core.ingest import docx as docx_mod

    d = D.Document()
    d.add_heading("第一章", level=1)
    t = d.add_table(rows=2, cols=4)
    for i, row in enumerate([["课程", "学分", "课程", "学分"], ["高数", "4.0", "英语", "3.0"]]):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    t2 = d.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "名称"; t2.cell(0, 1).text = "值"
    t2.cell(1, 0).text = "价格"; t2.cell(1, 1).text = "38"
    p = tmp_path / "t.docx"; d.save(str(p))
    monkeypatch.setattr(docx_mod, "_has_omml", lambda _doc: True)
    monkeypatch.setattr(docx_mod, "_pandoc_to_markdown", lambda _p: (
        "# 第一章\n\n|课程|学分|课程|学分|\n|---|---|---|---|\n|高数|4.0|英语|3.0|\n\n"
        "|名称|值|\n|---|---|\n|价格|38|\n"
    ))

    result = pl.ingest(p, ocr_name="off")
    assert len(result.tables) == 3                              # 子表切分：1 张 docx 表 → 2 张 + 第二张表
    assert "|课程|学分|课程|学分|" in result.parsed.full_text    # 数量对不上 → 原文一律保留，不按序错配
    assert "|名称|值|" in result.parsed.full_text
    assert "（表格见结构化数据）" not in result.parsed.full_text  # 不得用失真占位句覆盖原文


def test_structify_no_mislabeling_when_table_block_zero_yield(tmp_path):
    # 零产出的表（块文本保留 markdown 兜底原文）不得被别的表的摘要覆盖
    import docx as D
    from app.core.tables import table_summary

    d = D.Document()
    d.add_heading("第一章", level=1)
    t = d.add_table(rows=2, cols=2)
    for i, row in enumerate([["课程", "学分"], ["高数", "4.0"]]):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    t2 = d.add_table(rows=1, cols=2)   # 只有表头 → normalize_grid 零产出
    t2.cell(0, 0).text = "名称"; t2.cell(0, 1).text = "值"
    p = tmp_path / "t.docx"; d.save(str(p))

    result = pl.ingest(p, ocr_name="off")
    assert len(result.tables) == 1
    block = [b for b in result.parsed.blocks if b.type == "table"][1]
    assert block.text.startswith("| 名称 | 值 |")                       # 零产出表保留兜底原文，不置占位
    assert "（表格见结构化数据）" not in block.text                     # 别的表的摘要/占位都不得盖上来
    assert "| 名称 | 值 |" in result.parsed.full_text                   # full_text 中的原文同样保留
    # 该摘要只出现在它自己的位置，没有被复制到零产出表的位置
    assert result.parsed.full_text.count(table_summary(result.tables[0])) == 1


def test_ingest_syncs_table_text_out_of_full_text(tmp_path, monkeypatch):
    # Ruling B：structify 先于 clean_text——表格原文不得以「折叠后」形式残留在 full_text
    raw = "|A|B|\n|---|---|\n|1|2|"
    fake = ParsedDocument(
        kind="word", title=None, full_text="# 标题\n\n" + raw, headings=[],
        blocks=[Block(type="table", text="表格（1行×2列：A、B）", heading_path="标题")],
        engine_used="pandoc",
        tables=[ParsedTable(headers=["A", "B"], rows=[["1", "2"]])],
    )
    monkeypatch.setattr(proc.WordProcessor, "run", lambda *a, **k: fake)
    p = tmp_path / "x.docx"; p.write_bytes(b"dummy")  # 让 detect_format 能打开、按扩展名判 word
    result = pl.ingest(p, ocr_name="off")
    assert "|A|B|" not in result.parsed.full_text
    assert "表格（1行×2列" in result.parsed.full_text
