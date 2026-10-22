import io
import struct
import zlib

import docx as _docx
from app.core.ingest import docx


def test_parse_docx_headings_and_table(tmp_path):
    d = _docx.Document()
    d.add_heading("第一章", level=1)
    d.add_paragraph("正文")
    d.add_heading("1.1 小节", level=2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "名称"; t.cell(0, 1).text = "值"
    t.cell(1, 0).text = "价格"; t.cell(1, 1).text = "38"
    p = tmp_path / "t.docx"; d.save(str(p))

    parsed = docx.parse_docx(p)
    assert [h.level for h in parsed.headings] == [1, 2]
    # 表格改走直读：表内容进 parsed.tables，表块/正文用摘要代替原表文本（防表格原文重复进索引）
    assert "名称" in parsed.full_text and "价格" not in parsed.full_text
    assert parsed.tables[0].rows == [["价格", "38"]]
    assert any(b.type == "table" and b.text.startswith("表格（1行×2列：") for b in parsed.blocks)


def _tiny_png() -> bytes:
    """生成 1x1 透明 PNG 字节，供 add_picture 使用。"""
    def chunk(t: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + t
            + data
            + struct.pack(">I", zlib.crc32(t + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\x00\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", idat)
        + chunk(b"IEND", b"")
    )


def test_parse_docx_inline_image_block(tmp_path):
    d = _docx.Document()
    d.add_heading("带图章节", level=1)
    d.add_picture(io.BytesIO(_tiny_png()))
    p = tmp_path / "img.docx"; d.save(str(p))

    parsed = docx.parse_docx(p)
    assert any(b.type == "image" for b in parsed.blocks)
