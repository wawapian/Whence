import docx as D

from app.core import tables as tables_mod
from app.core.ingest import docx as docx_mod


def _make_docx(tmp_path, merge=None):
    d = D.Document()
    t = d.add_table(rows=3, cols=3)
    vals = [["课程", "性质", "学分"], ["高数", "必修", "4.0"], ["英语", "必修", "3.0"]]
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    if merge == "vertical":
        t.cell(1, 1).merge(t.cell(2, 1))
    if merge == "horizontal":
        t.cell(0, 0).merge(t.cell(0, 1))
    p = tmp_path / "t.docx"
    d.save(str(p))
    return p


def test_docx_table_grid_expands_vertical_merge(tmp_path):
    doc = D.Document(str(_make_docx(tmp_path, merge="vertical")))
    grid = docx_mod.docx_table_grid(doc.tables[0])
    assert len(grid) == 3 and all(len(r) == 3 for r in grid)
    assert grid[1][1] == grid[2][1]                       # 延续格与上一行同值
    ts = tables_mod.normalize_grid(grid)
    assert ts[0].rows == [["高数", "必修", "4.0"], ["英语", "必修", "3.0"]]  # "必修\n必修" → 去重


def test_docx_table_grid_expands_horizontal_merge(tmp_path):
    doc = D.Document(str(_make_docx(tmp_path, merge="horizontal")))
    grid = docx_mod.docx_table_grid(doc.tables[0])
    assert len(grid[0]) == 3 and grid[0][1] == ""
    assert grid[0][0] == "课程\n性质"                                  # 原始 cell.text
    assert tables_mod.normalize_grid(grid)[0].headers[0] == "课程性质"   # 清洗交给 normalize_grid


def test_docx_table_grid_horizontal_merge_keeps_all_paragraphs(tmp_path):
    """合并格原样保留各段：左格为空、标签在被并入格时不得读出空串，清洗后标签仍在。"""
    p = tmp_path / "m.docx"
    d = D.Document()
    t = d.add_table(rows=2, cols=3)
    vals = [["", "课程性质", "学分"], ["高数", "必修", "4.0"]]
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    t.cell(0, 0).merge(t.cell(0, 1))
    d.save(str(p))

    grid = docx_mod.docx_table_grid(D.Document(str(p)).tables[0])
    assert grid[0][1] == ""                                          # 其余列位置空
    assert tables_mod.normalize_grid(grid)[0].headers[0] == "课程性质"  # 被并入格的标签不丢


def test_parse_docx_collects_structured_tables(tmp_path):
    parsed = docx_mod.parse_docx(_make_docx(tmp_path))
    assert len(parsed.tables) == 1
    assert parsed.tables[0].headers == ["课程", "性质", "学分"]
    assert parsed.tables[0].rows == [["高数", "必修", "4.0"], ["英语", "必修", "3.0"]]
    blocks = [b for b in parsed.blocks if b.type == "table"]
    assert blocks and blocks[0].text.startswith("表格（2行×3列：")


def test_parse_docx_table_heading_is_heading_path(tmp_path):
    p = tmp_path / "h.docx"
    d = D.Document()
    d.add_heading("第一章", level=1)
    d.add_heading("1.1 小节", level=2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "课程"; t.cell(0, 1).text = "学分"
    t.cell(1, 0).text = "体育"; t.cell(1, 1).text = "1.0"
    d.save(str(p))

    parsed = docx_mod.parse_docx(p)
    assert parsed.tables[0].heading == "第一章 > 1.1 小节"


def test_parse_docx_table_block_falls_back_to_markdown_without_rows(tmp_path):
    p = tmp_path / "header_only.docx"
    d = D.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "课程"; t.cell(0, 1).text = "学分"
    d.save(str(p))

    parsed = docx_mod.parse_docx(p)
    assert parsed.tables == []                              # 只有表头 → 零产出
    block = [b for b in parsed.blocks if b.type == "table"][0]
    assert block.text.startswith("| 课程 | 学分 |")          # 退回旧 markdown 渲染，不丢内容


def test_parse_docx_pandoc_branch_fills_tables_and_summarizes_table_blocks(tmp_path, monkeypatch):
    p = _make_docx(tmp_path)
    monkeypatch.setattr(docx_mod, "_has_omml", lambda doc: True)
    monkeypatch.setattr(
        docx_mod, "_pandoc_to_markdown",
        lambda _p: "# 第一章\n\n|课程|性质|\n|---|---|\n|高数|必修|\n\n|a|b|\n|---|---|\n|1|2|\n",
    )

    parsed = docx_mod.parse_docx(p)
    assert parsed.engine_used == "pandoc"
    assert parsed.tables[0].headers == ["课程", "性质", "学分"]   # 从 doc.tables 直读
    assert parsed.tables[0].heading == "第一章"
    blocks = [b for b in parsed.blocks if b.type == "table"]
    assert blocks[0].text.startswith("表格（2行×3列：")            # 一一对应 → 摘要
    assert blocks[1].text == "|a|b|\n|---|---|\n|1|2|"            # 多出的表块：原文保留（不置占位）


def test_parse_docx_pandoc_branch_collects_tables_without_matching_block(tmp_path, monkeypatch):
    p = _make_docx(tmp_path)
    monkeypatch.setattr(docx_mod, "_has_omml", lambda doc: True)
    monkeypatch.setattr(docx_mod, "_pandoc_to_markdown", lambda _p: "# 第一章\n\n正文")

    parsed = docx_mod.parse_docx(p)
    assert [t.headers for t in parsed.tables] == [["课程", "性质", "学分"]]  # 多出的表仍收进
    assert parsed.tables[0].heading is None                                 # 无对应表块 → 无 heading
    assert not [b for b in parsed.blocks if b.type == "table"]
