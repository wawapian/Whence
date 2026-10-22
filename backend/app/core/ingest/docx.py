"""Word .docx 结构化解析：按 body 顺序提取标题层级、段落与表格。"""
import logging
import subprocess
from pathlib import Path

import docx as _docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.ingest import ocr
from app.core.ingest.markdown import parse_markdown
from app.core.ingest.models import Block, Heading, ImageRef, ParsedDocument
from app.core.tables import ParsedTable, normalize_grid, table_summary

logger = logging.getLogger(__name__)


def _heading_level(paragraph: Paragraph) -> int | None:
    """从段落样式名解析标题级别；非标题样式返回 None。"""
    name = (paragraph.style.name or "").strip()
    if name == "Title":
        return 1
    if "Heading" in name:
        digits = "".join(ch for ch in name if ch.isdigit())
        return int(digits) if digits else 1
    return None


def _has_image(paragraph: Paragraph) -> bool:
    """判断段落内是否包含内嵌图片（w:drawing 或 VML w:pict）。"""
    return bool(
        paragraph._p.xpath(".//w:drawing") or paragraph._p.xpath(".//w:pict")
    )


_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _image_blobs(paragraph: Paragraph) -> list[bytes]:
    """提取段落内嵌图片的字节（经 a:blip 的 r:embed 关系定位 target_part）。"""
    blobs: list[bytes] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(_R_EMBED)
        if not rid:
            continue
        part = paragraph.part.rels[rid].target_part
        blobs.append(part.blob)
    return blobs


def _table_to_markdown(table: Table) -> str:
    """逐格读取表格，渲染为 Markdown 表格文本（首行为表头）。"""
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    lines: list[str] = []
    for i, row in enumerate(rows):
        cells = row + [""] * (width - len(row))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + "---|" * width)
    return "\n".join(lines)


def docx_table_grid(table: Table) -> list[list[str]]:
    """读一张 docx 表为逻辑等宽网格（XML 合并标记展开，供 `normalize_grid` 消费）。

    - 横向合并：`row.cells` 会按 `gridSpan` 把同一个 `tc` 重复返回，按 `id(tc)` 去重后展开宽度——
      值置首个列位、其余列位置空。合并格的值取 `cell.text` 原样：python-docx / Word 合并会把被并入
      格的段落折进合并格（如 `"课程\\n性质"`），清洗与重复段落去重一律交给 `normalize_grid`。
    - 纵向合并：`row.cells` 已把 `vMerge` 延续格归到上方根格，其文本即根格内容
      （延续行沿用上一行同列值），无需另作处理。
    """
    grid: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            tc = cell._tc
            if id(tc) in seen:  # 横向合并：同一 tc 在行内重复出现
                continue
            seen.add(id(tc))
            span = tc.grid_span
            cells.append(cell.text)
            cells.extend([""] * (span - 1))
        grid.append(cells)
    width = max((len(row) for row in grid), default=0)
    return [row + [""] * (width - len(row)) for row in grid]


def _structured_tables(table: Table, heading: str | None) -> list[ParsedTable]:
    """直读一张 docx 表 → 结构化表（heading = 表所在 heading_path）。"""
    parsed = normalize_grid(docx_table_grid(table))
    for parsed_table in parsed:
        parsed_table.heading = heading
    return parsed


def _fill_pandoc_tables(parsed: ParsedDocument, doc) -> ParsedDocument:
    """pandoc 分支补结构化表：`doc.tables` 与 `type=="table"` 块按出现顺序一一对应。

    对应的表块文本替换为该表摘要（该表零产出则保留 pandoc 原文，不丢内容）；多出的表块保留
    pandoc 原文（不置占位——占位句宣称「见结构化数据」，但该表未必有结构化数据）；多出的表
    仍收进 `parsed.tables`。
    """
    table_blocks = [b for b in parsed.blocks if b.type == "table"]
    for block, table in zip(table_blocks, doc.tables):
        parsed_tables = _structured_tables(table, block.heading_path)
        parsed.tables.extend(parsed_tables)
        if parsed_tables:
            block.text = "\n".join(table_summary(t) for t in parsed_tables)
    for table in doc.tables[len(table_blocks):]:
        parsed.tables.extend(_structured_tables(table, None))
    return parsed


def _has_omml(doc) -> bool:
    """判断 docx 文档体是否包含 OMML 公式（m:oMath 元素）。"""
    return bool(doc.element.findall(".//" + qn("m:oMath")))


def _pandoc_to_markdown(path: Path) -> str:
    """用 pandoc 将 .docx 转为 markdown（公式渲染为 $...$/$$...$$ LaTeX）。"""
    result = subprocess.run(
        ["pandoc", str(path), "-t", "markdown"],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    if result.returncode != 0:
        raise RuntimeError(f"pandoc 转换失败: {result.stderr.strip()}")
    return result.stdout


def parse_docx(path: Path, ocr_name: str = "off") -> ParsedDocument:
    """解析 .docx 文件为 ParsedDocument，标题、段落、表格按文档顺序提取。

    表格走 python-docx 直读（合并标记展开）→ `parsed.tables`；表块文本替换为该表摘要，
    直读零产出时退回 `_table_to_markdown` 旧渲染（保底不丢内容）。
    """
    doc = _docx.Document(str(path))

    if _has_omml(doc):
        try:
            markdown_parsed = parse_markdown(
                _pandoc_to_markdown(path), kind="word", engine_used="pandoc"
            )
            return _fill_pandoc_tables(markdown_parsed, doc)
        except (FileNotFoundError, RuntimeError):
            pass  # pandoc 缺失/失败：回退 python-docx 结构化解析

    headings: list[Heading] = []
    blocks: list[Block] = []
    tables: list[ParsedTable] = []
    text_parts: list[str] = []
    images: list[ImageRef] = []
    title: str | None = None
    heading_stack: list[str] = []
    ocr_used = False
    engine = ocr.get_ocr_engine(ocr_name)

    def _heading_path() -> str | None:
        return " > ".join(heading_stack) if heading_stack else None

    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            level = _heading_level(item)
            if level is not None:
                if title is None and level == 1:
                    title = text
                headings.append(Heading(level=level, text=text))
                del heading_stack[level - 1:]
                heading_stack.append(text)
                blocks.append(
                    Block(type="heading", text=text, heading_path=_heading_path(), page=None)
                )
                text_parts.append(text)
            elif _has_image(item):
                try:
                    blobs = _image_blobs(item)
                except Exception:
                    blobs = []  # blob 提取失败：跳过该图，不影响整篇文档
                parts: list[str] = []
                heading_path = _heading_path()
                for blob in blobs:
                    ocr_text = ""
                    if engine is not None:
                        try:
                            ocr_text = engine.recognize(blob)
                        except Exception:
                            logger.warning("嵌入图片 OCR 失败（已跳过该图）", exc_info=True)
                            ocr_text = ""
                    if ocr_text:
                        parts.append(ocr_text)
                    images.append(
                        ImageRef(blob=blob, page=None, position=heading_path, ocr_text=ocr_text)
                    )
                image_text = "\n".join(parts).strip()
                if image_text:
                    ocr_used = True
                blocks.append(
                    Block(type="image", text=image_text, heading_path=heading_path, page=None)
                )
                if image_text:
                    text_parts.append(image_text)
            elif text:
                blocks.append(
                    Block(type="paragraph", text=text, heading_path=_heading_path(), page=None)
                )
                text_parts.append(text)
        elif isinstance(item, Table):
            heading_path = _heading_path()
            parsed_tables = _structured_tables(item, heading_path)
            tables.extend(parsed_tables)
            text = (
                "\n".join(table_summary(t) for t in parsed_tables)
                if parsed_tables
                else _table_to_markdown(item)
            )
            if text:
                blocks.append(
                    Block(type="table", text=text, heading_path=heading_path, page=None)
                )
                text_parts.append(text)

    return ParsedDocument(
        kind="word",
        title=title,
        full_text="\n".join(text_parts),
        headings=headings,
        blocks=blocks,
        engine_used="python-docx",
        ocr_used=ocr_used,
        needs_ocr=None,
        images=images,
        tables=tables,
    )
